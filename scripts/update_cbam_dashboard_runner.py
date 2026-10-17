#!/usr/bin/env python3
"""Compatibility runner for the live EEX annual auction workbook."""
from __future__ import annotations

from io import BytesIO
from openpyxl import load_workbook
import update_cbam_dashboard as base


def wide_find_header_row(ws):
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500), values_only=True), start=1
    ):
        labels = [base.norm(v) for v in row]
        date_i = next((i for i, x in enumerate(labels) if "auction date" in x), None)
        price_i = next((i for i, x in enumerate(labels) if "auction clearing price" in x), None)
        vol_i = next((i for i, x in enumerate(labels) if "auction volume" in x), None)
        if date_i is None or price_i is None or vol_i is None:
            continue
        mapping = {"date": date_i, "price": price_i, "volume": vol_i}
        status_i = next((i for i, x in enumerate(labels) if x == "status" or x.endswith(" status")), None)
        if status_i is not None:
            mapping["status"] = status_i
        contract_i = next((i for i, x in enumerate(labels) if x == "contract"), None)
        if contract_i is not None:
            mapping["contract"] = contract_i
        auction_i = next((i for i, x in enumerate(labels) if x == "auction name"), None)
        if auction_i is not None:
            mapping["auction_name"] = auction_i
        return row_idx, mapping
    return None


base.find_header_row = wide_find_header_row
_original_parse = base.parse_auction_xlsx


def debug_parse(content: bytes):
    rows = _original_parse(content)
    if rows:
        return rows
    wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    print("EEX workbook diagnostics:")
    for ws in wb.worksheets:
        print(f"SHEET {ws.title!r} rows={ws.max_row} cols={ws.max_column}")
        shown = 0
        for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [(i + 1, v) for i, v in enumerate(row) if v not in (None, "")]
            if vals:
                clipped = [(i, str(v)[:120]) for i, v in vals[:12]]
                print(f"ROW {ridx}: {clipped}")
                shown += 1
                if shown >= 35:
                    break
    return rows


base.parse_auction_xlsx = debug_parse
raise SystemExit(base.main())
